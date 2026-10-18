import ctypes
import logging
import sys

from dependency_check import dependency_message
from dependency_check import missing_dependencies
from logging_config import setup_logging


def main() -> None:
    """Start the SpaceWeekly desktop application."""
    missing = missing_dependencies()

    if missing:
        _show_dependency_error(dependency_message(missing))
        return

    from PySide6.QtGui import QIcon
    from PySide6.QtWidgets import QApplication
    from PySide6.QtWidgets import QDialog

    from resources import resource_path
    from services.workspace import needs_first_run_wizard
    from ui.first_run_wizard import FirstRunWizard

    app = QApplication(sys.argv)
    icon_path = resource_path("assets", "sspo.ico")

    if icon_path.exists():
        app.setWindowIcon(QIcon(str(icon_path)))
    else:
        logging.getLogger(__name__).warning("Icon not found %s", icon_path)

    if len(sys.argv) == 3 and sys.argv[1] == "--smoke-test":
        from pathlib import Path
        from services.workspace import initialize_workspace
        marker = Path(sys.argv[2])
        def checked(phase):
            with marker.with_suffix(".trace").open("a", encoding="utf-8") as trace:
                trace.write(phase + "\n")
        initialize_workspace(marker.parent / "workspace", "zh_CN")
        checked("workspace")
        from ui.settings_dialog import SettingsDialog
        from services.settings import default_settings
        from ui.theme import ThemeManager
        from config import app_version
        from ui.widgets.dashboard_panel import DashboardPanel
        from ui.video_cards_dialog import VideoCardsDialog
        from ui.video_intro_dialog import VideoIntroDialog
        from ui.video_download_dialog import VideoDownloadDialog
        from services.video_download import backend_command, _run_process
        from services.cancellation import no_cancel
        from models.video_download import VideoInfo, VideoFormat
        from qrcode import make as make_qr
        from qrcode.image.pil import PilImage
        from io import BytesIO
        from ui.prompt_studio_dialog import PromptStudioDialog
        from services.news_dates import publication_evidence
        from tools import video_cards, video_intro
        assert video_cards.FONT == video_cards.BUNDLED_FONT, 'Bundled text font was not selected'
        assert video_cards.NUMBER_FONT == video_cards.BUNDLED_FONT, 'Bundled number font was not selected'
        font = video_cards._load_font(video_cards.BUNDLED_FONT, 32)
        assert bytes(font.getmask('航')) != bytes(font.getmask('天')), 'Bundled Chinese glyphs failed'
        checked('bundled-fonts-selected')
        ThemeManager(app).apply("light")
        settings = default_settings()
        settings.deepseek_api_key = ""
        settings.web_search_api_key = ""
        dialog = SettingsDialog(settings)
        dialog.show()
        app.processEvents()
        dialog.close()
        checked("settings")
        dashboard = DashboardPanel(settings)
        cards = VideoCardsDialog(str(marker.parent))
        intro = VideoIntroDialog(str(marker.parent))
        logo = marker.parent / '测试 Logo.png'
        video_intro.default_logo().save(logo)
        intro.logo.setText(str(logo))
        intro._preview()
        assert intro.preview.pixmap() and not intro.preview.pixmap().isNull(), 'Opener preview with fallback font failed'
        checked('opener-font-fallback')
        intro._toggle_preview()
        assert intro._preview_timer.isActive(), 'Opener playback did not start'
        intro.preview_slider.setValue(120)
        intro._stop_preview()
        assert intro.preview_slider.value() == 120 and not intro._preview_timer.isActive()
        checked('opener-playback-and-seek')
        cards._add_launch()
        cards.launch_opacity.setValue(35)
        cards._prepare_preview()
        assert cards._preview_card['background_alpha'] == 89, 'Launch opacity failed'
        writing = PromptStudioDialog(settings)
        writing.type_list.setCurrentRow(2)
        assert writing.current_type == 'title' and writing.settings().title_prompt, 'Headline rule missing'
        writing.type_list.setCurrentRow(0)
        downloads = VideoDownloadDialog(marker.parent, [
            dict(video_type="launch", title="Smoke test", summary="Smoke mission",
                 published="2026-10-09T08:00:00+08:00")])
        downloads.rights.setChecked(True)
        assert downloads.demo.isEnabled(), "Packaged batch download UI failed"
        assert downloads.demo in downloads.more_menu.actions(), 'Batch action missing from More'
        assert downloads.open_file in downloads.more_menu.actions(), 'File action missing from More'
        checked('simplified-video-controls-and-headline-rule')
        formats = [VideoFormat("low", 360, "avc1"), VideoFormat("high", 1080, "avc1")]
        downloads._show_parsed_video(VideoInfo("smoke", "Smoke video",
            "https://www.bilibili.com/video/BV1xx411c7mD", formats=formats))
        assert downloads.formats.currentData() == formats[1], "Highest default quality failed"
        backend = Path(backend_command()[0])
        assert backend.name.lower() == "yt-dlp.exe" and backend.is_file(), "Download backend missing"
        backend_output = []
        checked("backend-start")
        _run_process([str(backend), "--version"], no_cancel, backend_output.append, timeout=30)
        assert backend_output and backend_output[-1] == "2026.08.19", "Packaged backend execution failed"
        checked("backend")
        import shutil
        import json
        ffmpeg, ffprobe = shutil.which("ffmpeg"), shutil.which("ffprobe")
        for tool in (ffmpeg, ffprobe):
            assert tool and Path(tool).parent == Path(sys._MEIPASS) / "ffmpeg", "Bundled media tool missing"
            _run_process([tool, "-version"], no_cancel, lambda line: None, timeout=30)
        checked("ffmpeg-start")
        parts, merged = marker.parent / "parts.mkv", marker.parent / "merged.mp4"
        _run_process([ffmpeg, "-v", "error", "-f", "lavfi", "-i", "testsrc2=s=64x64:r=10",
            "-f", "lavfi", "-i", "sine=frequency=440:sample_rate=44100", "-t", "0.3",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-threads", "1", str(parts)],
            no_cancel, lambda line: None, timeout=30)
        _run_process([ffmpeg, "-v", "error", "-i", str(parts), "-c", "copy", "-movflags", "+faststart", str(merged)],
            no_cancel, lambda line: None, timeout=30)
        probe = []
        _run_process([ffprobe, "-v", "error", "-show_streams", "-of", "json", str(merged)],
            no_cancel, probe.append, timeout=30)
        assert {stream["codec_name"] for stream in json.loads("\n".join(probe))["streams"]} == {"h264", "aac"}
        checked("mp4-encode-merge")
        for codec, pixels, extra in [("png", "rgba", []), ("prores_ks", "yuva444p10le", ["-profile:v", "4"])]:
            movie = marker.parent / (codec + ".mov")
            _run_process([ffmpeg, "-v", "error", "-f", "lavfi", "-i", "color=c=black@0:s=32x32:r=2,format=rgba",
                "-frames:v", "1", "-c:v", codec, "-pix_fmt", pixels, *extra, "-threads", "1", str(movie)],
                no_cancel, lambda line: None, timeout=30)
            probe = []
            _run_process([ffprobe, "-v", "error", "-show_streams", "-of", "json", str(movie)],
                no_cancel, probe.append, timeout=30)
            stream = json.loads("\n".join(probe))["streams"][0]
            assert stream["codec_name"] == ("png" if codec == "png" else "prores")
            assert stream["pix_fmt"] == "rgba" or stream["pix_fmt"].startswith("yuva")
        checked("alpha-mov-prores")
        for encoding in ('png', 'prores'):
            movie = video_cards.encode_card(cards._preview_card, marker.parent / encoding, 360,
                                            ffmpeg, ffprobe, encoding=encoding)
            assert movie.is_file(), 'Semi-transparent launch encoding failed'
        checked('launch-opacity-encoded')
        with BytesIO() as qr:
            make_qr("SpaceWeekly 3.0.2 package check", image_factory=PilImage).save(qr, format="PNG")
            assert qr.getvalue().startswith(bytes([137, 80, 78, 71])), "Packaged QR renderer failed"
        checked("qr")
        from docx import Document
        import requests
        import ssl
        word = Document()
        word.add_paragraph("SpaceWeekly packaged Word check")
        word_path = marker.parent / "smoke.docx"
        word.save(word_path)
        assert Document(word_path).paragraphs[0].text == "SpaceWeekly packaged Word check"
        checked("word")
        tls = ssl.create_default_context(cafile=requests.certs.where())
        assert tls.check_hostname and tls.verify_mode == ssl.CERT_REQUIRED
        checked("tls")
        from zoneinfo import ZoneInfo
        from datetime import datetime, timedelta
        for zone, month, hours in [
            ("Asia/Shanghai", 7, 8), ("UTC", 7, 0),
            ("America/New_York", 1, -5), ("America/New_York", 7, -4),
        ]:
            assert datetime(2025, month, 1, tzinfo=ZoneInfo(zone)).utcoffset() == timedelta(hours=hours)
        checked("timezones")
        assert not QIcon(str(resource_path("assets", "ui-check.svg"))).pixmap(16, 16).isNull()
        checked("svg")
        downloads.show()
        app.processEvents()
        downloads.close()
        checked("downloads")
        assert writing.current_type == "filter", "Writing rules failed to initialize"
        evidence = publication_evidence(
            '<meta name="PubDate" content="2026-09-30"><article>'
            + "A fictional mission announces its first science images and begins science operations."
            + "</article>", "https://www.cas.cn/test")
        assert evidence["published"] == "2026-09-30", "Packaged date parser failed"
        frame = video_intro.render_frame(video_intro.prepare("2026-01-01", end_date="2026-01-07"), 3)
        assert frame.size == (1920, 1080), "Packaged opener renderer failed"
        opener = video_intro.generate(marker.parent / 'encoded-opener', **intro.options())
        assert opener.is_file() and opener.stat().st_size > 0, 'Packaged full opener encoding failed'
        checked("opener-encoded-420frames")
        checked("opener")
        writing.close()
        dashboard.close()
        cards.close()
        intro.close()
        checked("dialogs-closed")
        from PySide6.QtCore import QTimer
        QTimer.singleShot(0, app.quit)
        app.exec()
        checked("qt-shutdown")
        marker.write_text(app_version(), encoding="utf-8")
        return

    if needs_first_run_wizard():
        wizard = FirstRunWizard()

        if wizard.exec() != QDialog.DialogCode.Accepted:
            return

    setup_logging()
    from ui.main_window import MainWindow

    window = MainWindow()
    window.show()
    sys.exit(app.exec())


def _show_dependency_error(message: str) -> None:
    if sys.platform.startswith("win"):
        ctypes.windll.user32.MessageBoxW(None, message, "SpaceWeekly", 0x10)
        return

    sys.stderr.write(message + "\n")


if __name__ == "__main__":
    main()
