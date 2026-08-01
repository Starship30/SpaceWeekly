import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

_TEST_APPDATA = Path(tempfile.gettempdir()) / "spaceweekly-v242-tests"
_TEST_APPDATA.mkdir(parents=True, exist_ok=True)
os.environ["APPDATA"] = str(_TEST_APPDATA)

from ai.json_utils import AIResponseParseError
from ai.json_utils import parse_json_object
from ai.json_utils import request_json
from ai.prompts import build_content_prompt
from ai.prompts import build_importance_prompt
from models.ai_summary import AISummary
from models.article import Article
from models.news import News
from models.news_assessment import NewsAssessment
from services.news_categories import category_counts
from services.news_categories import category_labels
from services.news_categories import category_limit
from services.news_categories import normalize_categories
from utils.output_files import reserved_output
from utils.text import clean_ai_text


def sample_article() -> Article:
    return Article(
        news=News(
            source="Test",
            title="Mission Update",
            published="Thu, 02 Jul 2026 10:00:00 GMT",
            summary="RSS summary",
            url="https://example.com/news",
        ),
        body="Article body",
        parser="Test",
    )


class OutputFileTests(unittest.TestCase):
    def test_numbering_skips_existing_files(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            (output / "SpaceWeekly_2026-07-26_001.docx").write_bytes(b"old")

            with reserved_output(
                output,
                ".docx",
                report_date=date(2026, 7, 26),
            ) as (final_path, temp_path):
                self.assertEqual(
                    final_path.name,
                    "SpaceWeekly_2026-07-26_002.docx",
                )
                temp_path.write_bytes(b"new")

            self.assertEqual(final_path.read_bytes(), b"new")

    def test_failed_export_removes_reservation_and_temp_file(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)

            with self.assertRaises(RuntimeError):
                with reserved_output(output, ".md") as (final_path, temp_path):
                    temp_path.write_text("partial", encoding="utf-8")
                    raise RuntimeError("cancelled")

            self.assertFalse(final_path.exists())
            self.assertFalse(temp_path.exists())

    def test_word_and_markdown_exports_are_numbered_and_readable(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is unavailable")

        import exporters.markdown as markdown
        import exporters.word as word

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            previous_word = word.OUTPUT_DIR
            previous_markdown = markdown.OUTPUT_DIR

            try:
                word.OUTPUT_DIR = output
                markdown.OUTPUT_DIR = output
                word_one = word.export_word([])
                word_two = word.export_word([])
                markdown_one = markdown.export_markdown([])
                markdown_two = markdown.export_markdown([])
                Document(word_one)
            finally:
                word.OUTPUT_DIR = previous_word
                markdown.OUTPUT_DIR = previous_markdown

            self.assertTrue(word_one.name.endswith("_001.docx"))
            self.assertTrue(word_two.name.endswith("_002.docx"))
            self.assertTrue(markdown_one.name.endswith("_001.md"))
            self.assertTrue(markdown_two.name.endswith("_002.md"))


class AIJsonTests(unittest.TestCase):
    def test_strips_markdown_and_repairs_common_errors(self):
        parsed = parse_json_object('```json\n{"summary": "ok",}\n```')
        self.assertEqual(parsed["summary"], "ok")
        truncated = parse_json_object('{"summary": "truncated')
        self.assertEqual(truncated["summary"], "truncated")

    def test_retries_once_after_parse_failure(self):
        provider = type("Config", (), {"model": "test-model"})()
        responses = ["not-json", '{"importance":"High","reason":"important"}']

        with patch("ai.providers.complete", side_effect=responses) as complete:
            result = request_json(
                "prompt",
                provider,
                module_name="test",
            )

        self.assertEqual(result["importance"], "High")
        self.assertEqual(complete.call_count, 2)

    def test_schema_validator_participates_in_retry(self):
        provider = type("Config", (), {"model": "test-model"})()
        responses = [
            '{"importance":"High","reason":"ok","summary":"forbidden"}',
            '{"importance":"High","reason":"ok"}',
        ]

        def validator(data):
            if set(data) != {"importance", "reason"}:
                raise AIResponseParseError("extra fields")

        with patch("ai.providers.complete", side_effect=responses) as complete:
            result = request_json(
                "prompt",
                provider,
                module_name="test",
                validator=validator,
            )

        self.assertEqual(set(result), {"importance", "reason"})
        self.assertEqual(complete.call_count, 2)

    def test_invalid_category_retries_once_with_same_stable_schema(self):
        from ai.deepseek import _validate_content

        provider = type("Config", (), {"model": "test-model"})()
        responses = [
            json.dumps(
                {
                    "summary": "有效摘要",
                    "keywords": ["关键词一", "关键词二", "关键词三"],
                    "category": ["临时分类"],
                },
                ensure_ascii=False,
            ),
            json.dumps(
                {
                    "summary": "有效摘要",
                    "keywords": ["关键词一", "关键词二", "关键词三"],
                    "category": ["科学发现"],
                },
                ensure_ascii=False,
            ),
        ]

        with patch("ai.providers.complete", side_effect=responses) as complete:
            result = request_json(
                "prompt",
                provider,
                module_name="category-test",
                validator=lambda value: _validate_content(value, "zh_CN"),
            )

        self.assertEqual(result["category"], ["科学发现"])
        self.assertEqual(complete.call_count, 2)


class WorkflowPromptTests(unittest.TestCase):
    def test_screening_prompt_has_only_minimal_schema(self):
        prompt = build_importance_prompt(sample_article(), "zh_CN")
        self.assertIn('"importance"', prompt)
        self.assertIn('"reason"', prompt)
        self.assertIn("禁止返回摘要", prompt)

    def test_content_prompt_has_controlled_category_schema(self):
        prompt = build_content_prompt(sample_article(), "en_US")
        self.assertIn('"summary"', prompt)
        self.assertIn('"keywords"', prompt)
        self.assertIn('"category"', prompt)
        self.assertIn("Launch Missions", prompt)
        self.assertIn("zero to 3 items", prompt)
        self.assertIn("Do not return title", prompt)

    def test_text_cleanup_removes_empty_quotes_and_bad_spaces(self):
        self.assertEqual(
            clean_ai_text('发现了 “ ” 新的 宇宙现象'),
            "发现了新的宇宙现象",
        )


class SettingsMigrationTests(unittest.TestCase):
    def test_old_report_style_migrates_to_filter_level(self):
        import services.settings as settings_module

        legacy = {
            "deepseek_model": "legacy-model",
            "report_style": "curated",
            "pipeline_translation_enabled": True,
            "category_prompt": "obsolete",
        }

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.json"
            path.write_text(json.dumps(legacy), encoding="utf-8")
            previous = settings_module.CONFIG_PATH

            try:
                settings_module.CONFIG_PATH = path
                settings = settings_module.load_settings()
                settings_module.save_settings(settings)
            finally:
                settings_module.CONFIG_PATH = previous

            saved = json.loads(path.read_text(encoding="utf-8"))

        self.assertEqual(settings.news_filter_level, "High_only")
        self.assertEqual(settings.deepseek_model, "legacy-model")
        self.assertNotIn("pipeline_translation_enabled", saved)
        self.assertNotIn("category_prompt", saved)


class CategorySystemTests(unittest.TestCase):
    def test_representative_news_uses_canonical_categories(self):
        from ai.deepseek import _summary_from_data
        from ai.deepseek import _validate_content

        cases = [
            (
                "发射类新闻",
                ["发射任务", "运载火箭", "商业航天"],
            ),
            (
                "科学发现新闻",
                ["天文观测", "科学发现", "深空探索"],
            ),
            (
                "卫星新闻",
                ["卫星与通信", "航天器"],
            ),
            (
                "Space.com 娱乐观星新闻",
                ["天文观测"],
            ),
        ]

        for name, categories in cases:
            with self.subTest(name=name):
                data = {
                    "summary": "这是一段符合当前语言设置的中文测试摘要内容。",
                    "keywords": ["航天", "任务", "进展"],
                    "category": categories,
                }
                _validate_content(data, "zh_CN")
                summary = _summary_from_data(data, "zh_CN")
                self.assertEqual(summary.categories, categories)

    def test_unknown_duplicate_and_excess_categories_are_rejected(self):
        from ai.deepseek import _validate_content

        base = {
            "summary": "有效摘要",
            "keywords": ["关键词一", "关键词二", "关键词三"],
        }
        invalid_categories = [
            ["先进商业火箭发射技术进展"],
            ["发射任务", "发射任务"],
            ["发射任务", "运载火箭", "商业航天", "航天技术"],
        ]

        for categories in invalid_categories:
            with self.subTest(categories=categories):
                with self.assertRaises(AIResponseParseError):
                    _validate_content(
                        {**base, "category": categories},
                        "zh_CN",
                    )

    def test_category_config_normalization_and_statistics_interface(self):
        self.assertEqual(category_limit(), 3)
        self.assertEqual(len(category_labels("zh_CN")), 12)
        self.assertEqual(
            normalize_categories(
                ["launch_missions", "Launch Vehicles", "商业航天"],
                "zh_CN",
            ),
            ["发射任务", "运载火箭", "商业航天"],
        )
        summaries = {
            "one": AISummary("摘要", ["关键词"], ["发射任务", "商业航天"]),
            "two": AISummary("摘要", ["关键词"], ["发射任务"]),
        }
        self.assertEqual(
            category_counts(summaries),
            {"发射任务": 2, "商业航天": 1},
        )

    def test_old_summary_without_categories_defaults_to_empty_list(self):
        self.assertEqual(AISummary("旧摘要", ["旧关键词"]).categories, [])

    def test_word_exports_category_and_importance_once(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is unavailable")

        import exporters.word as word

        article = sample_article()
        summary = AISummary(
            "测试摘要",
            ["星舰", "SpaceX", "火箭"],
            ["发射任务", "商业航天"],
        )
        assessment = NewsAssessment("High", "重要测试任务")

        with tempfile.TemporaryDirectory() as directory:
            previous = word.OUTPUT_DIR

            try:
                word.OUTPUT_DIR = Path(directory)
                path = word.export_word(
                    [article],
                    {article.news.url: summary},
                    assessments={article.news.url: assessment},
                    launch_range="disabled",
                )
                document = Document(path)
            finally:
                word.OUTPUT_DIR = previous

        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertEqual(text.count("分类："), 1)
        self.assertEqual(text.count("重要程度："), 1)
        self.assertIn("发射任务、商业航天", text)


if __name__ == "__main__":
    unittest.main()
